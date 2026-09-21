"""The downloadable analysis workbook.

This is the artefact the NPCU assembles by hand each quarter, and the one the
Bank and the states actually read. It has to say exactly what the dashboard
says, because the two go out together and a reader who finds them disagreeing
will trust neither.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.services import analysis_model
from app.services.reporting import workbook
from tests.test_analysis_model import (  # noqa: F401  (fixtures)
    _indicator,
    _report,
    _submission,
    quarter,
)


@pytest.fixture()
def book(db, quarter):  # noqa: F811
    model = analysis_model.build(db, "2026-Q1")
    stream = io.BytesIO()
    workbook.build(model).save(stream)
    stream.seek(0)
    return load_workbook(stream), model


def test_the_sheets_are_the_ones_the_npcu_workbook_has(book):
    wb, model = book
    assert wb.sheetnames[:2] == ["Cover", "National Summary"]
    assert "Data Quality Flags" in wb.sheetnames
    assert "Full Data Table" in wb.sheetnames
    from app.services.analysis_model import COMPONENT_SHORT

    for component in model.components:
        assert COMPONENT_SHORT[component] in wb.sheetnames


def test_a_sheet_name_is_never_cut_mid_word(book):
    """Excel caps a name at 31 characters, which turned the full component
    title into "Component 1 — Creating Safe and"."""
    wb, _ = book
    for name in wb.sheetnames:
        assert len(name) <= 31
        assert not name.endswith(("and", "the", "of", "—"))


def test_a_component_sheet_carries_every_state_as_a_column(book):
    wb, model = book
    sheet = wb["PDO"]
    header = [sheet.cell(row=4, column=c).value for c in range(1, sheet.max_column + 1)]

    assert header[:3] == ["Code", "Indicator", "Type"]
    assert header[3:3 + len(model.states)] == model.states
    assert header[-3:] == ["National Total", "Target", "% of Target"]


def test_an_unreported_figure_stays_blank(book):
    """A blank and a zero are different conversations with a state."""
    wb, model = book
    sheet = wb["PDO"]
    lagos = 4 + model.states.index("Lagos")
    row = next(
        r for r in range(5, sheet.max_row + 1)
        if sheet.cell(row=r, column=1).value == "PDO-90"
    )
    assert sheet.cell(row=row, column=lagos).value is None


def test_a_flagged_figure_is_shaded_and_still_carries_its_value(book):
    """Shading, not exclusion. Holding the doubtful figures out is what made
    this platform's totals disagree with the NPCU's published report."""
    wb, model = book
    sheet = wb["Component 1"]
    kano = 4 + model.states.index("Kano")
    row = next(
        r for r in range(5, sheet.max_row + 1)
        if sheet.cell(row=r, column=1).value == "C1-90"
    )
    cell = sheet.cell(row=row, column=kano)

    assert cell.value == 120.0, "the flagged figure keeps its value"
    assert cell.fill.fgColor.rgb not in (None, "00000000"), "and is shaded"

    national = sheet.cell(row=row, column=4 + len(model.states))
    assert national.value == 120.0, "and still counts towards the national total"


def test_a_boolean_reads_yes_or_no(db, quarter):  # noqa: F811
    flag = _indicator(db, "C3-91", component="C3", unit="BOOLEAN", method="COUNT_YES",
                      name="Policy adopted?")
    from sqlalchemy import select

    from app.models import Submission
    from app.services import reference

    period = reference.get_period_by_code(db, "2026-Q1")
    submission = db.scalar(
        select(Submission).where(
            Submission.is_current.is_(True), Submission.period_id == period.id
        )
    )
    _report(db, submission, flag, 1.0)
    db.commit()

    model = analysis_model.build(db, "2026-Q1")
    stream = io.BytesIO()
    workbook.build(model).save(stream)
    stream.seek(0)
    sheet = load_workbook(stream)["Component 3"]

    row = next(
        r for r in range(5, sheet.max_row + 1)
        if sheet.cell(row=r, column=1).value == "C3-91"
    )
    values = {sheet.cell(row=row, column=c).value for c in range(4, 4 + len(model.states))}
    assert "Yes" in values
    assert 1 not in values and 1.0 not in values


def test_the_flags_sheet_keeps_each_state_s_own_wording(book):
    wb, _ = book
    sheet = wb["Data Quality Flags"]
    header = [sheet.cell(row=4, column=c).value for c in range(1, 6)]
    assert header == ["Severity", "Code", "Indicator", "States", "Finding"]


def test_the_full_table_skips_unreported_figures(book):
    """A row of zero for a state that filed nothing is a fabricated figure."""
    wb, model = book
    sheet = wb["Full Data Table"]
    states = {
        sheet.cell(row=r, column=5).value for r in range(2, sheet.max_row + 1)
    }
    codes = [
        sheet.cell(row=r, column=2).value for r in range(2, sheet.max_row + 1)
    ]
    assert "Lagos" not in states, "Lagos reported nothing this period"
    assert "PDO-90" in codes


def test_the_endpoint_serves_the_workbook(client, npcu_headers, quarter):  # noqa: F811
    response = client.get(
        "/api/v1/reports/analysis-workbook",
        headers=npcu_headers,
        params={"period": "2026-Q1"},
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "2026-Q1" in response.headers["content-disposition"]
    assert load_workbook(io.BytesIO(response.content)).sheetnames[0] == "Cover"
