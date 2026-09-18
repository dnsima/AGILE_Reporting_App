"""The query resolution screen, from the API's side.

The rule the whole workflow exists to enforce: a state proposes a correction
with evidence, the NPCU approves it, and only then does the stored figure
change. These tests hold the endpoints to that, and to the access rules that
make it mean something -- a state answers for its own figures and nobody
else's, and nobody clears their own response into the national total.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone
from itertools import count

import pytest
from openpyxl import load_workbook

from app.core.enums import QueryResolution, QueryStatus, SubmissionStatus
from app.models import DataQuery, Indicator, IndicatorValue, Submission
from app.services import corrections, reference
from app.services import queries as query_service
from app.services.validation import run_validation

_numbers = count(700)


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------
def _indicator(db, code="C1.0-02", **kwargs):
    indicator = Indicator(
        code=code,
        number=next(_numbers),
        name=kwargs.pop("name", "Constructed or rehabilitated SSS classrooms"),
        unit=kwargs.pop("unit", "NUMBER"),
        aggregation_method="SUM",
        direction="INCREASE",
        **kwargs,
    )
    db.add(indicator)
    db.flush()
    return indicator


def _submission(db, state="KN", period="2026-Q1"):
    state_row = reference.get_state_by_code(db, state)
    period_row = reference.get_period_by_code(db, period)
    version = 1
    for existing in db.query(Submission).filter_by(
        state_id=state_row.id, period_id=period_row.id
    ):
        existing.is_current = False
        version = max(version, existing.version + 1)
    submission = Submission(
        state_id=state_row.id,
        period_id=period_row.id,
        version=version,
        status=str(SubmissionStatus.APPROVED),
        is_current=True,
        uploaded_at=datetime.now(timezone.utc),
    )
    db.add(submission)
    db.flush()
    return submission


def _value(db, submission, indicator, value):
    row = IndicatorValue(
        submission_id=submission.id,
        indicator_id=indicator.id,
        value=value,
        original_value=value,
        disaggregation={},
    )
    db.add(row)
    db.flush()
    return row


@pytest.fixture()
def flagged(db):
    """Kebbi's real case, on Kano: a cumulative figure that falls between quarters."""
    indicator = _indicator(db, is_cumulative=True)
    earlier = _submission(db, period="2025-Q4")
    _value(db, earlier, indicator, 1577.0)
    later = _submission(db, period="2026-Q1")
    value = _value(db, later, indicator, 342.0)
    run_validation(db, later)
    raised = query_service.raise_queries(db, later)
    query = next(q for q in raised if q.rule_code == "CON-002")
    db.commit()
    return {
        "indicator": indicator,
        "earlier": earlier,
        "later": later,
        "value": value,
        "query": query,
        "query_id": query.id,
    }


def _get(client, headers, path, **params):
    response = client.get(path, headers=headers, params=params)
    assert response.status_code == 200, response.text
    return response.json()


# --------------------------------------------------------------------------
# Worklists
# --------------------------------------------------------------------------
class TestWorklist:
    def test_the_state_sees_the_finding_in_the_words_the_npcu_saw(
        self, client, state_headers, flagged
    ):
        rows = _get(client, state_headers, "/api/v1/queries", period="2026-Q1")
        row = next(r for r in rows if r["id"] == flagged["query_id"])

        assert row["reference"] == f"Q-{flagged['query_id']}"
        assert row["rule_code"] == "CON-002"
        assert "fell from 1,577 to 342" in row["title"]
        assert row["reported_value"] == 342.0
        assert row["is_open"] is True
        assert row["due_date"] is not None

    def test_a_held_figure_says_so(self, client, npcu_headers, flagged):
        rows = _get(client, npcu_headers, "/api/v1/queries", period="2026-Q1")
        row = next(r for r in rows if r["id"] == flagged["query_id"])
        assert row["is_quarantined"] is True
        assert row["current_value"] == 342.0

    def test_a_state_sees_only_its_own_queries(self, client, state_headers, db, flagged):
        other = _submission(db, state="KD", period="2026-Q1")
        _value(db, other, flagged["indicator"], 10.0)
        earlier = _submission(db, state="KD", period="2025-Q4")
        _value(db, earlier, flagged["indicator"], 900.0)
        run_validation(db, other)
        query_service.raise_queries(db, other)
        db.commit()

        rows = _get(client, state_headers, "/api/v1/queries", period="2026-Q1")
        assert {row["state_code"] for row in rows} == {"KN"}

        denied = client.get("/api/v1/queries?state=KD", headers=state_headers)
        assert denied.status_code == 403

    def test_a_state_cannot_open_another_states_query(
        self, client, state_headers, db, flagged
    ):
        other = _submission(db, state="KD", period="2026-Q1")
        _value(db, other, flagged["indicator"], 10.0)
        earlier = _submission(db, state="KD", period="2025-Q4")
        _value(db, earlier, flagged["indicator"], 900.0)
        run_validation(db, other)
        raised = query_service.raise_queries(db, other)
        db.commit()

        response = client.get(f"/api/v1/queries/{raised[0].id}", headers=state_headers)
        assert response.status_code == 403

    def test_overdue_is_visible_and_filterable(self, client, npcu_headers, db, flagged):
        query = db.get(DataQuery, flagged["query_id"])
        query.due_date = date.today() - timedelta(days=9)
        db.commit()

        rows = _get(client, npcu_headers, "/api/v1/queries", overdue_only=True)
        row = next(r for r in rows if r["id"] == flagged["query_id"])
        assert row["is_overdue"] is True
        assert row["days_overdue"] == 9

    def test_summary_counts_the_period(self, client, npcu_headers, flagged):
        body = _get(client, npcu_headers, "/api/v1/queries/summary", period="2026-Q1")
        assert body["period_code"] == "2026-Q1"
        assert body["open"] >= 1
        assert body["awaiting_review"] == 0
        assert body["restated"] == 0

    def test_the_summary_route_is_not_shadowed_by_the_detail_route(
        self, client, npcu_headers, flagged
    ):
        """`/queries/summary` and `/queries/{id}` are both one segment."""
        body = _get(client, npcu_headers, "/api/v1/queries/summary")
        assert "total" in body


# --------------------------------------------------------------------------
# The state's half
# --------------------------------------------------------------------------
class TestResponding:
    def test_confirming_the_figure_leaves_it_alone(self, client, state_headers, db, flagged):
        body = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=state_headers,
            json={
                "narrative": "The figure is right; the Q4 return double-counted.",
                "evidence_summary": "Contractor certificates, June site visit",
            },
        )
        assert body.status_code == 201, body.text
        detail = body.json()
        assert detail["status"] == str(QueryStatus.RESPONDED)
        assert detail["responses"][0]["proposed_value"] is None
        assert detail["responses"][0]["responder"] == "Kano PIU"

        db.expire_all()
        assert db.get(IndicatorValue, flagged["value"].id).value == 342.0

    def test_a_proposed_correction_does_not_change_the_figure_yet(
        self, client, state_headers, db, flagged
    ):
        """Nothing moves until the NPCU accepts. That is the whole control."""
        response = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=state_headers,
            json={
                "narrative": "The Q4 figure was overstated; evidence supports 342.",
                "proposed_value": 342.0,
                "restates_period_code": "2025-Q4",
            },
        )
        assert response.status_code == 201, response.text

        db.expire_all()
        earlier_value = db.scalar(
            db.query(IndicatorValue)
            .filter_by(submission_id=flagged["earlier"].id)
            .statement
        )
        assert earlier_value.value == 1577.0
        assert db.get(DataQuery, flagged["query_id"]).status == str(QueryStatus.RESPONDED)

    def test_a_response_must_say_something(self, client, state_headers, flagged):
        response = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=state_headers,
            json={"narrative": "   "},
        )
        assert response.status_code == 422

    def test_a_state_cannot_respond_for_another_state(
        self, client, state_headers, db, flagged
    ):
        other = _submission(db, state="KD", period="2026-Q1")
        _value(db, other, flagged["indicator"], 10.0)
        earlier = _submission(db, state="KD", period="2025-Q4")
        _value(db, earlier, flagged["indicator"], 900.0)
        run_validation(db, other)
        raised = query_service.raise_queries(db, other)
        db.commit()

        response = client.post(
            f"/api/v1/queries/{raised[0].id}/responses",
            headers=state_headers,
            json={"narrative": "Not mine to answer."},
        )
        assert response.status_code == 403

    def test_evidence_is_attached_and_fingerprinted(self, client, state_headers, flagged):
        created = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=state_headers,
            json={"narrative": "Evidence attached.", "proposed_value": 342.0},
        )
        response_id = created.json()["responses"][0]["id"]

        uploaded = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses/{response_id}/evidence",
            headers=state_headers,
            files={"file": ("certificates.pdf", b"%PDF-1.4 certificates", "application/pdf")},
            data={"description": "Signed contractor certificates"},
        )
        assert uploaded.status_code == 201, uploaded.text
        evidence = uploaded.json()["responses"][0]["evidence"]
        assert len(evidence) == 1
        assert evidence[0]["filename"] == "certificates.pdf"
        assert len(evidence[0]["content_hash"]) == 64
        assert evidence[0]["description"] == "Signed contractor certificates"

    def test_an_empty_evidence_file_is_refused(self, client, state_headers, flagged):
        created = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=state_headers,
            json={"narrative": "Evidence to follow."},
        )
        response_id = created.json()["responses"][0]["id"]
        uploaded = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses/{response_id}/evidence",
            headers=state_headers,
            files={"file": ("empty.pdf", b"", "application/pdf")},
        )
        assert uploaded.status_code == 422


# --------------------------------------------------------------------------
# The NPCU's half
# --------------------------------------------------------------------------
class TestReview:
    def _responded(self, client, state_headers, query_id, **payload):
        body = {"narrative": "Evidence supports the restatement."}
        body.update(payload)
        response = client.post(
            f"/api/v1/queries/{query_id}/responses", headers=state_headers, json=body
        )
        assert response.status_code == 201, response.text
        return response.json()

    def test_a_state_cannot_clear_its_own_response(
        self, client, state_headers, flagged
    ):
        self._responded(client, state_headers, flagged["query_id"], proposed_value=342.0)
        response = client.post(
            f"/api/v1/queries/{flagged['query_id']}/accept",
            headers=state_headers,
            json={"note": "Approving my own figure."},
        )
        assert response.status_code == 403

    def test_accepting_a_restatement_changes_the_figure_and_keeps_the_original(
        self, client, state_headers, npcu_headers, db, flagged
    ):
        self._responded(
            client,
            state_headers,
            flagged["query_id"],
            proposed_value=342.0,
            restates_period_code="2025-Q4",
        )
        accepted = client.post(
            f"/api/v1/queries/{flagged['query_id']}/accept",
            headers=npcu_headers,
            json={"note": "Certificates checked."},
        )
        assert accepted.status_code == 200, accepted.text
        detail = accepted.json()
        assert detail["status"] == str(QueryStatus.ACCEPTED)
        assert detail["resolution"] == str(QueryResolution.RESTATED)
        assert detail["is_open"] is False

        db.expire_all()
        earlier = db.scalar(
            db.query(IndicatorValue).filter_by(submission_id=flagged["earlier"].id).statement
        )
        assert earlier.value == 342.0
        assert earlier.original_value == 1577.0  # never overwritten
        assert earlier.is_restated is True

    def test_accepting_releases_the_held_figure(
        self, client, state_headers, npcu_headers, db, flagged
    ):
        self._responded(client, state_headers, flagged["query_id"])
        client.post(
            f"/api/v1/queries/{flagged['query_id']}/accept",
            headers=npcu_headers,
            json={"note": "Confirmed."},
        )
        db.expire_all()
        assert db.get(IndicatorValue, flagged["value"].id).is_valid is True

    def test_rejecting_returns_it_to_the_state(
        self, client, state_headers, npcu_headers, flagged
    ):
        self._responded(client, state_headers, flagged["query_id"])
        rejected = client.post(
            f"/api/v1/queries/{flagged['query_id']}/reject",
            headers=npcu_headers,
            json={"reason": "The certificates do not cover the period."},
        )
        assert rejected.status_code == 200, rejected.text
        detail = rejected.json()
        assert detail["status"] == str(QueryStatus.REJECTED)
        assert detail["is_open"] is True  # still the state's work
        assert detail["responses"][0]["review_outcome"] == "REJECTED"

    def test_referring_for_verification_keeps_it_open(
        self, client, npcu_headers, flagged
    ):
        referred = client.post(
            f"/api/v1/queries/{flagged['query_id']}/verification",
            headers=npcu_headers,
            json={"reason": "Check the classrooms on the next supervision visit."},
        )
        assert referred.status_code == 200, referred.text
        assert referred.json()["verification_required"] is True
        assert referred.json()["is_open"] is True

        rows = _get(client, npcu_headers, "/api/v1/queries", verification_only=True)
        assert flagged["query_id"] in {row["id"] for row in rows}

    def test_withdrawing_releases_the_figure(
        self, client, npcu_headers, db, flagged
    ):
        withdrawn = client.post(
            f"/api/v1/queries/{flagged['query_id']}/withdraw",
            headers=npcu_headers,
            json={"reason": "The rule was retuned; this was never a real finding."},
        )
        assert withdrawn.status_code == 200, withdrawn.text
        assert withdrawn.json()["status"] == str(QueryStatus.WITHDRAWN)

        db.expire_all()
        assert db.get(IndicatorValue, flagged["value"].id).is_valid is True

    def test_accepting_with_no_response_is_refused(self, client, npcu_headers, flagged):
        response = client.post(
            f"/api/v1/queries/{flagged['query_id']}/accept",
            headers=npcu_headers,
            json={"note": "Nothing to go on."},
        )
        assert response.status_code == 409

    def test_a_viewer_can_read_but_not_act(self, client, viewer_headers, flagged):
        assert _get(client, viewer_headers, "/api/v1/queries", period="2026-Q1")
        blocked = client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=viewer_headers,
            json={"narrative": "Not my place."},
        )
        assert blocked.status_code == 403


# --------------------------------------------------------------------------
# Correction sheets
# --------------------------------------------------------------------------
class TestCorrectionSheets:
    def test_the_sheet_carries_only_this_states_flagged_figures(
        self, client, state_headers, db, flagged
    ):
        response = client.get(
            "/api/v1/queries/sheets/KN?period=2026-Q1", headers=state_headers
        )
        assert response.status_code == 200, response.text
        assert "AGILE_corrections_KN_2026-Q1.xlsx" in response.headers["content-disposition"]

        sheet = load_workbook(io.BytesIO(response.content))["Corrections"]
        refs = {
            row[0]
            for row in sheet.iter_rows(min_row=6, values_only=True)
            if row and row[0]
        }
        open_for_kano = {
            f"Q-{query.id}"
            for query in query_service.open_queries(
                db,
                state_id=reference.get_state_by_code(db, "KN").id,
                period_id=reference.get_period_by_code(db, "2026-Q1").id,
            )
        }
        assert f"Q-{flagged['query_id']}" in refs
        # Only this state's flagged figures, and every one of them: a figure
        # nobody queried cannot be changed by this route.
        assert refs == open_for_kano

    def test_a_state_cannot_pull_another_states_sheet(self, client, state_headers, flagged):
        response = client.get(
            "/api/v1/queries/sheets/KD?period=2026-Q1", headers=state_headers
        )
        assert response.status_code == 403

    def test_a_completed_sheet_becomes_responses(
        self, client, state_headers, db, flagged
    ):
        state = reference.get_state_by_code(db, "KN")
        period = reference.get_period_by_code(db, "2026-Q1")
        payload = corrections.build_correction_sheet(db, state, period)

        workbook = load_workbook(io.BytesIO(payload))
        sheet = workbook["Corrections"]
        row = next(
            index
            for index, values in enumerate(sheet.iter_rows(min_row=6, values_only=True), start=6)
            if values and values[0] == f"Q-{flagged['query_id']}"
        )
        sheet.cell(row=row, column=7, value="Corrected")
        sheet.cell(row=row, column=8, value=342)
        sheet.cell(row=row, column=9, value="Contractor certificates")
        sheet.cell(row=row, column=10, value="The Q4 figure was overstated.")
        buffer = io.BytesIO()
        workbook.save(buffer)

        response = client.post(
            "/api/v1/queries/sheets/KN",
            headers=state_headers,
            files={
                "file": (
                    "corrections.xlsx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"period_code": "2026-Q1"},
        )
        assert response.status_code == 201, response.text
        body = response.json()
        assert body["applied"] == 1
        assert body["skipped"] == 0

        db.expire_all()
        query = db.get(DataQuery, flagged["query_id"])
        assert query.status == str(QueryStatus.RESPONDED)
        assert query.responses[-1].proposed_value == 342.0
        # And still nothing has moved until the NPCU accepts.
        assert db.get(IndicatorValue, flagged["value"].id).value == 342.0

    def test_a_sheet_with_nothing_filled_in_is_refused(
        self, client, state_headers, db, flagged
    ):
        state = reference.get_state_by_code(db, "KN")
        period = reference.get_period_by_code(db, "2026-Q1")
        payload = corrections.build_correction_sheet(db, state, period)

        response = client.post(
            "/api/v1/queries/sheets/KN",
            headers=state_headers,
            files={"file": ("untouched.xlsx", payload, "application/vnd.ms-excel")},
            data={"period_code": "2026-Q1"},
        )
        assert response.status_code == 422
        assert "at least one" in response.json()["error"]["message"]


class TestScoping:
    def test_a_state_is_counted_against_its_own_queries_not_the_nation(
        self, client, state_headers, npcu_headers, db, flagged
    ):
        """106 open nationally against a state's own 21 is not a usable number."""
        other = _submission(db, state="KD", period="2026-Q1")
        _value(db, other, flagged["indicator"], 10.0)
        earlier = _submission(db, state="KD", period="2025-Q4")
        _value(db, earlier, flagged["indicator"], 900.0)
        run_validation(db, other)
        query_service.raise_queries(db, other)
        db.commit()

        national = _get(client, npcu_headers, "/api/v1/queries/summary", period="2026-Q1")
        mine = _get(client, state_headers, "/api/v1/queries/summary", period="2026-Q1")

        assert national["open"] > mine["open"]
        assert mine["open"] == len(
            _get(client, state_headers, "/api/v1/queries", period="2026-Q1")
        )
        assert mine["states_with_open_queries"] == 1

    def test_a_state_cannot_ask_for_another_states_counts(self, client, state_headers, flagged):
        response = client.get(
            "/api/v1/queries/summary?period=2026-Q1&state=KD", headers=state_headers
        )
        assert response.status_code == 403


class TestSharedFigures:
    """Two rules can flag one figure; settling one does not settle the other."""

    def _both(self, db, flagged):
        extra = DataQuery(
            submission_id=flagged["later"].id,
            state_id=flagged["later"].state_id,
            period_id=flagged["later"].period_id,
            indicator_id=flagged["indicator"].id,
            rule_code="INT-005",
            dimension="INTEGRITY",
            severity="ERROR",
            title="C1.0-02 reports 342 but its parts sum to 237.",
            reported_value=342.0,
            status=str(QueryStatus.OPEN),
            due_date=date.today() + timedelta(days=14),
        )
        db.add(extra)
        db.flush()
        db.commit()
        return extra

    def test_the_detail_names_what_else_is_holding_the_figure(
        self, client, npcu_headers, db, flagged
    ):
        extra = self._both(db, flagged)
        body = _get(client, npcu_headers, f"/api/v1/queries/{flagged['query_id']}")
        assert body["held_by"] == [f"Q-{extra.id}"]
        assert body["is_quarantined"] is True

    def test_settling_one_query_leaves_the_figure_held_by_the_other(
        self, client, state_headers, npcu_headers, db, flagged
    ):
        extra = self._both(db, flagged)
        client.post(
            f"/api/v1/queries/{flagged['query_id']}/responses",
            headers=state_headers,
            json={"narrative": "The figure is right."},
        )
        client.post(
            f"/api/v1/queries/{flagged['query_id']}/accept",
            headers=npcu_headers,
            json={"note": "Confirmed."},
        )
        db.expire_all()
        assert db.get(IndicatorValue, flagged["value"].id).is_valid is False

        # Settle the second, and the figure returns to the national total.
        client.post(
            f"/api/v1/queries/{extra.id}/withdraw",
            headers=npcu_headers,
            json={"reason": "Raised against a superseded composite definition."},
        )
        db.expire_all()
        assert db.get(IndicatorValue, flagged["value"].id).is_valid is True
