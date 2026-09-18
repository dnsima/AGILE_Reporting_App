"""The reporting template states fill in.

The template is the contract between a state and the platform, so these tests
hold it to the four things that make the round trip safe: every row carries its
code, a state is only asked for what it implements, only its own approved
targets are shown, and what a state types into the entry column is what the
platform reads back out.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.core.enums import IndicatorUnit
from app.models import (
    Indicator,
    ReportingPeriod,
    State,
    StateSubcomponent,
    Subcomponent,
    Target,
)
from app.services import reference
from app.services.ingestion import build_reporting_template, map_rows, parse_upload
from app.services.ingestion.template import reporting_basis

HEADER_ROW = 6


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def _sheet(payload: bytes, name: str = "Reporting"):
    return load_workbook(io.BytesIO(payload))[name]


def _headers(sheet) -> list[str]:
    return [sheet.cell(row=HEADER_ROW, column=c).value for c in range(1, sheet.max_column + 1)]


def _entry_column(sheet, period: ReportingPeriod) -> int:
    return _headers(sheet).index(f"Reported value ({period.code})") + 1


def _rows(sheet) -> dict[str, dict[str, object]]:
    """``{indicator code: {header: cell value}}`` for the indicator rows only."""
    headers = _headers(sheet)
    body: dict[str, dict[str, object]] = {}
    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        code = sheet.cell(row=row, column=1).value
        # Component banners span the sheet and leave every other cell blank.
        if not code or sheet.cell(row=row, column=3).value is None:
            continue
        body[str(code)] = {
            header: sheet.cell(row=row, column=index).value
            for index, header in enumerate(headers, start=1)
        }
    return body


def _split_components(db) -> tuple[Subcomponent, Subcomponent]:
    """Put two indicators under 1.1 and two under 1.2, as the real catalogue does."""
    first = Subcomponent(code="C1.1", name="New builds", sort_order=1)
    second = Subcomponent(code="C1.2", name="Existing schools", sort_order=2)
    db.add_all([first, second])
    db.flush()

    indicators = db.query(Indicator).order_by(Indicator.number).all()
    indicators[0].subcomponent_id = first.id
    indicators[1].subcomponent_id = first.id
    indicators[2].subcomponent_id = second.id
    indicators[3].subcomponent_id = second.id
    db.flush()
    return first, second


@pytest.fixture()
def period(db) -> ReportingPeriod:
    return reference.get_period_by_code(db, "2026-Q1")


@pytest.fixture()
def kano(db) -> State:
    return reference.get_state_by_code(db, "KN")


# --------------------------------------------------------------------------
# Reporting basis
# --------------------------------------------------------------------------
def test_basis_is_stated_per_indicator_not_per_column(db):
    """One column header cannot say 'cumulative' and 'snapshot' at once."""
    cumulative, rate, snapshot, boolean = db.query(Indicator).order_by(Indicator.number).all()
    rate.unit = str(IndicatorUnit.PERCENT)
    snapshot.is_cumulative = False
    boolean.unit = str(IndicatorUnit.BOOLEAN)

    assert "Cumulative total to date" in reporting_basis(cumulative)
    assert "between 0 and 100" in reporting_basis(rate)
    assert "Do not enter a count" in reporting_basis(rate)
    assert "snapshot" in reporting_basis(snapshot)
    assert reporting_basis(boolean) == "Enter Yes or No."


def test_composite_basis_names_its_parts(db):
    composite = db.query(Indicator).order_by(Indicator.number).first()
    composite.composite_of = ["KPI-002", "KPI-003"]
    assert reporting_basis(composite) == (
        "Cumulative total to date. Must equal KPI-002 + KPI-003."
    )
    # A part the state does not implement is dropped from the arithmetic it is
    # asked to satisfy, rather than named in a sum it cannot complete.
    assert reporting_basis(composite, ["KPI-003"]) == (
        "Cumulative total to date. Must equal KPI-003."
    )


# --------------------------------------------------------------------------
# Shape
# --------------------------------------------------------------------------
def test_every_row_carries_its_code(db, kano, period):
    sheet = _sheet(build_reporting_template(db, kano, period))
    rows = _rows(sheet)
    assert set(rows) == {"KPI-001", "KPI-002", "KPI-003", "KPI-004"}
    assert _headers(sheet)[0] == "Code"


def test_cover_block_names_the_state_and_period(db, kano, period):
    sheet = _sheet(build_reporting_template(db, kano, period))
    assert "Kano" in sheet["A1"].value
    assert period.label in sheet["A1"].value
    assert "Results framework" in sheet["A1"].value
    assert f"Period: {period.code}" in sheet["A3"].value
    assert period.due_date.isoformat() in sheet["A2"].value


def test_monthly_period_issues_the_tracker_layout(db, kano):
    monthly = reference.ensure_period(db, "MONTHLY", 2026, 2)
    db.flush()
    sheet = _sheet(build_reporting_template(db, kano, monthly))
    assert "Performance tracker" in sheet["A1"].value
    # Same catalogue, same codes: that is what makes the monthly-to-quarterly
    # reconciliation a code join rather than a match on indicator names.
    assert set(_rows(sheet)) == {"KPI-001", "KPI-002", "KPI-003", "KPI-004"}


def test_workbook_carries_guidance_and_reference_tabs(db, kano, period):
    workbook = load_workbook(io.BytesIO(build_reporting_template(db, kano, period)))
    assert workbook.sheetnames == ["Reporting", "How to complete", "Indicator reference"]
    reference_rows = list(workbook["Indicator reference"].iter_rows(values_only=True))
    assert reference_rows[0][0] == "Code"
    assert {row[0] for row in reference_rows[1:]} == {
        "KPI-001", "KPI-002", "KPI-003", "KPI-004"
    }


def _first_indicator_row(sheet) -> int:
    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        if sheet.cell(row=row, column=3).value is not None:
            return row
    raise AssertionError("the template has no indicator rows")


def test_entry_cells_are_unlocked_and_the_rest_is_not(db, kano, period):
    sheet = _sheet(build_reporting_template(db, kano, period))
    entry = _entry_column(sheet, period)
    row = _first_indicator_row(sheet)
    assert sheet.protection.sheet is True
    assert sheet.cell(row=row, column=entry).protection.locked is False
    assert sheet.cell(row=row, column=entry + 1).protection.locked is False  # Data source
    assert sheet.cell(row=row, column=1).protection.locked is True           # Code
    assert sheet.cell(row=row, column=6).protection.locked is True           # Target


# --------------------------------------------------------------------------
# Applicability
# --------------------------------------------------------------------------
def test_rows_a_state_does_not_implement_are_locked_out(db, kano, period):
    first, second = _split_components(db)
    db.add_all([
        StateSubcomponent(state_id=kano.id, subcomponent_id=first.id, implements=False),
        StateSubcomponent(state_id=kano.id, subcomponent_id=second.id, implements=True),
    ])
    db.flush()

    sheet = _sheet(build_reporting_template(db, kano, period))
    rows = _rows(sheet)
    entry = _headers(sheet)[_entry_column(sheet, period) - 1]

    assert rows["KPI-001"][entry] == "n/a"
    assert rows["KPI-002"][entry] == "n/a"
    assert rows["KPI-003"][entry] is None
    assert rows["KPI-004"][entry] is None

    # The row says why it is blank, so the blank is never read as a data gap.
    assert "does not implement C1.1" in rows["KPI-001"]["How to report this figure"]
    assert "Leave blank" in rows["KPI-001"]["How to report this figure"]


def test_a_state_with_no_matrix_is_served_the_full_catalogue(db, kano, period):
    _split_components(db)
    sheet = _sheet(build_reporting_template(db, kano, period))
    entry = _headers(sheet)[_entry_column(sheet, period) - 1]
    assert all(row[entry] is None for row in _rows(sheet).values())


def test_composite_arithmetic_drops_a_locked_out_part(db, kano, period):
    first, second = _split_components(db)
    composite, part_one, part_two = db.query(Indicator).order_by(Indicator.number).all()[:3]
    composite.composite_of = [part_one.code, part_two.code]
    db.add_all([
        StateSubcomponent(state_id=kano.id, subcomponent_id=first.id, implements=True),
        StateSubcomponent(state_id=kano.id, subcomponent_id=second.id, implements=False),
    ])
    db.flush()

    basis = _rows(_sheet(build_reporting_template(db, kano, period)))
    # KPI-002 is under 1.1 with the composite; KPI-003 is under the locked 1.2.
    assert basis[composite.code]["How to report this figure"] == (
        f"Cumulative total to date. Must equal {part_one.code}."
    )


# --------------------------------------------------------------------------
# Targets
# --------------------------------------------------------------------------
def test_only_the_states_own_approved_target_is_served(db, kano, period):
    indicators = db.query(Indicator).order_by(Indicator.number).all()
    targets = {
        target.indicator_id: target
        for target in db.query(Target).filter_by(state_id=kano.id, period_id=period.id)
    }
    targets[indicators[1].id].status = "DRAFT"
    db.query(Target).filter_by(
        state_id=kano.id, period_id=period.id, indicator_id=indicators[2].id
    ).delete()
    # A national target is not this state's target, and must not stand in for one.
    db.add(
        Target(
            indicator_id=indicators[2].id, period_id=period.id, level="NATIONAL",
            state_id=None, target_value=99999.0, status="APPROVED",
        )
    )
    db.flush()

    rows = _rows(_sheet(build_reporting_template(db, kano, period)))
    assert rows["KPI-001"]["Target"] == 1000.0     # approved for this state
    assert rows["KPI-002"]["Target"] == "—"        # approved by nobody yet
    assert rows["KPI-003"]["Target"] == "—"        # national only


# --------------------------------------------------------------------------
# Prior periods
# --------------------------------------------------------------------------
def test_prior_columns_show_what_the_state_reported(db, client, state_headers, kano, period):
    earlier = reference.get_period_by_code(db, "2025-Q4")
    response = client.post(
        "/api/v1/ingestion/submissions?auto_approve=false",
        headers=state_headers,
        json={
            "state_code": "KN",
            "period_code": earlier.code,
            "values": [{"indicator_code": "KPI-001", "value": 412}],
        },
    )
    assert response.status_code == 201, response.text

    db.expire_all()
    sheet = _sheet(build_reporting_template(db, kano, period))
    column = f"{earlier.label} — as reported"
    assert column in _headers(sheet)
    rows = _rows(sheet)
    assert rows["KPI-001"][column] == 412
    assert rows["KPI-002"][column] == "—"


def test_a_period_the_state_never_reported_is_not_shown(db, kano, period):
    sheet = _sheet(build_reporting_template(db, kano, period))
    assert not [header for header in _headers(sheet) if "as reported" in str(header)]


# --------------------------------------------------------------------------
# Round trip
# --------------------------------------------------------------------------
def _fill(payload: bytes, period: ReportingPeriod, figures: dict[str, float]) -> bytes:
    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook["Reporting"]
    entry = _entry_column(sheet, period)
    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        code = sheet.cell(row=row, column=1).value
        if code in figures:
            sheet.cell(row=row, column=entry, value=figures[code])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_a_filled_template_maps_back_onto_its_own_indicators(db, kano, period):
    figures = {"KPI-001": 1200.0, "KPI-002": 64.5, "KPI-003": 38.0, "KPI-004": 91.0}
    filled = _fill(build_reporting_template(db, kano, period), period, figures)

    sheet = parse_upload(filled, "filled.xlsx")
    result = map_rows(sheet.headers, sheet.rows, reference.active_indicators(db))

    assert result.layout == "long"
    assert result.unmapped_rows == 0
    assert result.unmatched_indicators == []
    assert {value.indicator.code: value.value for value in result.values} == figures


def test_component_banners_are_not_charged_as_unmapped_rows(db, kano, period):
    """The template's own formatting must not cost the state an integrity finding."""
    payload = build_reporting_template(db, kano, period)
    sheet = parse_upload(payload, "empty.xlsx")
    assert any("section heading" in warning for warning in sheet.warnings)

    result = map_rows(sheet.headers, sheet.rows, reference.active_indicators(db))
    assert result.unmapped_rows == 0
    assert result.mapped_rows == 4


def test_history_columns_are_never_read_as_this_periods_figure(db, client, state_headers, kano, period):
    """The reference columns sit beside the entry column; only one is the answer."""
    earlier = reference.get_period_by_code(db, "2025-Q4")
    client.post(
        "/api/v1/ingestion/submissions",
        headers=state_headers,
        json={
            "state_code": "KN",
            "period_code": earlier.code,
            "values": [{"indicator_code": "KPI-001", "value": 412}],
        },
    )
    db.expire_all()

    filled = _fill(build_reporting_template(db, kano, period), period, {"KPI-001": 500.0})
    sheet = parse_upload(filled, "filled.xlsx")
    result = map_rows(sheet.headers, sheet.rows, reference.active_indicators(db))

    value_headers = [
        mapping.source_header
        for mapping in result.column_mappings
        if mapping.mapped_field == "value"
    ]
    assert value_headers == [f"Reported value ({period.code})"]
    assert {v.indicator.code: v.value for v in result.values}["KPI-001"] == 500.0


def test_a_locked_row_submits_nothing(db, kano, period):
    first, second = _split_components(db)
    db.add_all([
        StateSubcomponent(state_id=kano.id, subcomponent_id=first.id, implements=False),
        StateSubcomponent(state_id=kano.id, subcomponent_id=second.id, implements=True),
    ])
    db.flush()

    payload = build_reporting_template(db, kano, period)
    sheet = parse_upload(payload, "untouched.xlsx")
    result = map_rows(sheet.headers, sheet.rows, reference.active_indicators(db))

    values = {value.indicator.code: value.value for value in result.values}
    assert values["KPI-001"] is None  # 'n/a' is a blank, not a zero
    assert values["KPI-002"] is None


# --------------------------------------------------------------------------
# Endpoint
# --------------------------------------------------------------------------
def test_endpoint_serves_the_callers_own_state_template(client, state_headers):
    response = client.get("/api/v1/ingestion/template?period=2026-Q1", headers=state_headers)
    assert response.status_code == 200, response.text
    assert "AGILE_results_framework_KN_2026-Q1.xlsx" in response.headers["content-disposition"]
    assert "Kano" in _sheet(response.content)["A1"].value


def test_endpoint_needs_a_state_to_build_against(client, npcu_headers):
    response = client.get("/api/v1/ingestion/template?period=2026-Q1", headers=npcu_headers)
    assert response.status_code == 422
    assert "Name the state" in response.json()["error"]["message"]


def test_a_state_cannot_pull_another_states_template(client, state_headers):
    response = client.get(
        "/api/v1/ingestion/template?state=KD&period=2026-Q1", headers=state_headers
    )
    assert response.status_code == 403


def test_endpoint_defaults_to_the_latest_quarter(client, npcu_headers):
    response = client.get("/api/v1/ingestion/template?state=KD", headers=npcu_headers)
    assert response.status_code == 200, response.text
    assert "AGILE_results_framework_KD_" in response.headers["content-disposition"]


# --------------------------------------------------------------------------
# The two streams agree on what a figure means
# --------------------------------------------------------------------------
def test_the_template_and_the_reconciliation_read_a_figure_the_same_way(db, kano, period):
    """A template that says "snapshot" while reconciliation sums is a trap."""
    from app.core.enums import TimeBasis
    from app.services import reconciliation

    workbook = load_workbook(io.BytesIO(build_reporting_template(db, kano, period)))
    rows = list(workbook["Indicator reference"].iter_rows(values_only=True))
    headers = list(rows[0])
    basis_col = headers.index("Basis")
    rollup_col = headers.index("Monthly rolls up as")

    catalogue = {row.code: row for row in db.query(Indicator)}
    for row in rows[1:]:
        indicator = catalogue[row[0]]
        sums = reconciliation.time_basis(indicator) is TimeBasis.SUM
        assert sums is ("added together" in row[rollup_col])
        # The instruction the state reads must not contradict the arithmetic.
        assert sums is ("this period only" in row[basis_col])


def test_a_flow_indicator_is_told_not_to_carry_figures_forward(db):
    from app.core.enums import TimeBasis

    indicator = db.query(Indicator).order_by(Indicator.number).first()
    indicator.is_cumulative = False
    indicator.time_basis = str(TimeBasis.SUM)
    assert reporting_basis(indicator) == (
        "Count for this period only. Do not carry anything forward from earlier periods."
    )
