"""Correction sheets: only flagged figures, with evidence, under approval."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from itertools import count

import pytest
from openpyxl import load_workbook

from app.core.enums import QueryStatus, SubmissionStatus
from app.core.errors import IngestionError, ValidationError
from app.models import Indicator, IndicatorValue, Submission
from app.services import corrections, queries, reference
from app.services.validation import run_validation

_numbers = count(700)


def _indicator(db, code, **kwargs):
    indicator = Indicator(
        code=code,
        number=next(_numbers),
        name=kwargs.pop("name", f"Indicator {code}"),
        unit="NUMBER",
        aggregation_method="SUM",
        direction="INCREASE",
        **kwargs,
    )
    db.add(indicator)
    db.flush()
    return indicator


def _submission(db, state="KN", period="2026-Q1"):
    state_row = reference.get_state_by_code(db, state)
    period_row = reference.get_period_by_code(db, period)
    previous = (
        db.query(Submission)
        .filter_by(state_id=state_row.id, period_id=period_row.id)
        .order_by(Submission.version.desc())
        .first()
    )
    if previous is not None:
        previous.is_current = False
    submission = Submission(
        state_id=state_row.id,
        period_id=period_row.id,
        version=(previous.version + 1) if previous else 1,
        status=str(SubmissionStatus.APPROVED),
        is_current=True,
        uploaded_at=datetime.now(timezone.utc),
    )
    db.add(submission)
    db.flush()
    return submission


def _value(db, submission, indicator, value):
    row = IndicatorValue(
        submission_id=submission.id,
        indicator_id=indicator.id,
        value=value,
        original_value=value,
        disaggregation={},
    )
    db.add(row)
    db.flush()
    return row


@pytest.fixture()
def flagged(db):
    """A state with one flagged figure: a cumulative total that fell."""
    indicator = _indicator(db, "C1.0-02", is_cumulative=True)
    earlier = _submission(db, period="2025-Q4")
    _value(db, earlier, indicator, 1577.0)
    later = _submission(db, period="2026-Q1")
    value = _value(db, later, indicator, 342.0)
    run_validation(db, later)
    raised = queries.raise_queries(db, later)
    query = next(q for q in raised if q.rule_code == "CON-002")
    state = reference.get_state_by_code(db, "KN")
    period = reference.get_period_by_code(db, "2026-Q1")
    return {
        "indicator": indicator, "value": value, "query": query,
        "state": state, "period": period, "submission": later,
    }


def _row_for(sheet, query_id):
    """Find the sheet row carrying a given query, rather than assuming an order."""
    for row in range(6, sheet.max_row + 1):
        if sheet.cell(row, 1).value == f"Q-{query_id}":
            return row
    raise AssertionError(f"Q-{query_id} is not on the sheet")


def _fill(sheet_bytes, query_id, **cells):
    """Fill the row for one query and return the workbook bytes."""
    workbook = load_workbook(io.BytesIO(sheet_bytes))
    sheet = workbook["Corrections"]
    row = _row_for(sheet, query_id)
    columns = {"response": 7, "corrected": 8, "evidence": 9, "explanation": 10}
    for name, value in cells.items():
        sheet.cell(row, columns[name], value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestBuilding:
    def test_the_sheet_carries_only_flagged_figures(self, db, flagged):
        """Sound figures never appear, so this route cannot change them."""
        sound = _indicator(db, "C9.0-01")
        _value(db, flagged["submission"], sound, 100.0)

        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        sheet = load_workbook(io.BytesIO(content))["Corrections"]

        codes = {sheet.cell(r, 2).value for r in range(6, sheet.max_row + 1)}
        open_now = queries.open_queries(
            db, state_id=flagged["state"].id, period_id=flagged["period"].id
        )
        assert "C1.0-02" in codes
        assert "C9.0-01" not in codes
        assert len([r for r in range(6, sheet.max_row + 1) if sheet.cell(r, 1).value]) == len(open_now)

    def test_the_row_shows_what_was_reported_and_what_was_flagged(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        sheet = load_workbook(io.BytesIO(content))["Corrections"]

        row = _row_for(sheet, flagged["query"].id)
        assert sheet.cell(row, 2).value == "C1.0-02"
        assert sheet.cell(row, 5).value == 342.0
        assert "fell from 1577 to 342" in sheet.cell(row, 6).value

    def test_entry_columns_are_left_blank_for_the_state(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        sheet = load_workbook(io.BytesIO(content))["Corrections"]
        row = _row_for(sheet, flagged["query"].id)
        assert [sheet.cell(row, column).value for column in (7, 8, 9, 10)] == [None] * 4

    def test_a_state_with_nothing_flagged_gets_no_sheet(self, db, flagged):
        clean = reference.get_state_by_code(db, "LA")
        with pytest.raises(ValidationError, match="nothing to correct"):
            corrections.build_correction_sheet(db, clean, flagged["period"])


class TestReturning:
    def test_a_correction_becomes_a_proposed_figure(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        filled = _fill(
            content, flagged["query"].id, response="CORRECTED", corrected=300,
            evidence="Handover certificates, 12 LGAs",
            explanation="Q4 figure overstated; only 300 had completion evidence.",
        )

        result = corrections.ingest_correction_sheet(
            db, filled, state=flagged["state"], period=flagged["period"]
        )
        assert result.applied == 1

        query = flagged["query"]
        assert query.status == QueryStatus.RESPONDED
        assert query.responses[-1].proposed_value == 300.0
        assert "Handover certificates" in query.responses[-1].evidence_summary

    def test_a_confirmation_proposes_no_change(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        filled = _fill(
            content, flagged["query"].id, response="CONFIRMED", evidence="SPIU monitoring report",
            explanation="Verified against the monitoring report; the figure stands.",
        )

        corrections.ingest_correction_sheet(
            db, filled, state=flagged["state"], period=flagged["period"]
        )
        assert flagged["query"].responses[-1].proposed_value is None

    def test_nothing_is_written_until_the_npcu_accepts(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        filled = _fill(
            content, flagged["query"].id, response="CORRECTED", corrected=300,
            evidence="Register", explanation="Overstated.",
        )
        corrections.ingest_correction_sheet(
            db, filled, state=flagged["state"], period=flagged["period"]
        )

        assert flagged["value"].value == 342.0, "figure changed before review"

        queries.accept(db, flagged["query"], note="Accepted.")
        assert flagged["value"].value == 300.0
        assert flagged["value"].original_value == 342.0

    def test_a_correction_without_a_figure_is_refused(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        filled = _fill(content, flagged["query"].id, response="CORRECTED", explanation="This looks wrong.")

        result = corrections.ingest_correction_sheet(
            db, filled, state=flagged["state"], period=flagged["period"]
        )
        assert result.applied == 0
        assert result.skipped == 1
        assert "no corrected figure" in result.warnings[0]
        assert flagged["query"].status == QueryStatus.OPEN

    def test_a_response_without_an_explanation_is_refused(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        filled = _fill(content, flagged["query"].id, response="CONFIRMED", evidence="A report")

        result = corrections.ingest_correction_sheet(
            db, filled, state=flagged["state"], period=flagged["period"]
        )
        assert result.applied == 0
        assert "explanation is required" in result.warnings[0]

    def test_untouched_rows_are_ignored(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        with pytest.raises(IngestionError, match="No completed rows"):
            corrections.ingest_correction_sheet(
                db, content, state=flagged["state"], period=flagged["period"]
            )


class TestControls:
    def test_a_sheet_cannot_answer_another_state_s_query(self, db, flagged):
        other_indicator = _indicator(db, "C1.0-09", is_cumulative=True)
        earlier = _submission(db, state="GO", period="2025-Q4")
        _value(db, earlier, other_indicator, 900.0)
        later = _submission(db, state="GO", period="2026-Q1")
        _value(db, later, other_indicator, 100.0)
        run_validation(db, later)
        foreign = queries.raise_queries(db, later)[0]

        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["Corrections"]
        row = _row_for(sheet, flagged["query"].id)
        sheet.cell(row, 1, f"Q-{foreign.id}")
        sheet.cell(row, 7, "CORRECTED")
        sheet.cell(row, 8, 99999)
        sheet.cell(row, 10, "Changing another state's figure.")
        buffer = io.BytesIO()
        workbook.save(buffer)

        result = corrections.ingest_correction_sheet(
            db, buffer.getvalue(), state=flagged["state"], period=flagged["period"]
        )
        assert result.applied == 0
        assert "does not belong to" in result.warnings[0]
        assert foreign.status == QueryStatus.OPEN

    def test_a_settled_query_cannot_be_reopened_by_a_sheet(self, db, flagged):
        content = corrections.build_correction_sheet(db, flagged["state"], flagged["period"])
        queries.respond(db, flagged["query"], narrative="Checked.")
        queries.accept(db, flagged["query"])

        filled = _fill(
            content, flagged["query"].id, response="CORRECTED", corrected=1, explanation="Changing my mind."
        )
        result = corrections.ingest_correction_sheet(
            db, filled, state=flagged["state"], period=flagged["period"]
        )
        assert result.applied == 0
        assert "already ACCEPTED" in result.warnings[0]

    def test_a_file_that_is_not_a_correction_sheet_is_refused(self, db, flagged):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["something", "else"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        with pytest.raises(IngestionError, match="not look like a correction sheet"):
            corrections.ingest_correction_sheet(
                db, buffer.getvalue(), state=flagged["state"], period=flagged["period"]
            )
