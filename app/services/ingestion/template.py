"""Generates the reporting templates states fill in.

Both layouts -- the quarterly results framework and the monthly performance
tracker -- come from the one indicator catalogue, so the two streams carry
identical codes and wording and monthly-to-quarterly reconciliation is an exact
code join rather than a match on indicator names.

The template is generated per state per period, which is what lets it design out
the faults the NPCU's consolidation notes record:

* every row carries its **indicator code**, so a reworded label cannot silently
  drop a figure (the current tracker matches on name alone);
* the state is **bound at upload**, never read from the file, so a cloned
  template cannot submit one state's figures under another's name;
* **targets are served by the platform**, not typed by the state, and only once
  the governance body has approved them;
* each row states **the basis its figure must be on** -- cumulative to date, a
  month-end snapshot, a rate, Yes/No -- instead of leaving a grid of monthly
  cells that different states fill three different ways;
* rows for sub-components a state does not implement are **locked and marked
  not applicable**, so a blank there is never mistaken for a missing figure.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import AggregationMethod, IndicatorUnit, PeriodType, TargetStatus
from app.models import (
    Indicator,
    ReportingPeriod,
    State,
    StateSubcomponent,
    Submission,
    Target,
)
from app.services import reference

HEADER_FILL = PatternFill("solid", fgColor="0F3D61")
BANNER_FILL = PatternFill("solid", fgColor="DCE6F1")
LOCKED_FILL = PatternFill("solid", fgColor="F2F2F2")
ENTRY_FILL = PatternFill("solid", fgColor="FFF9E6")
BLOCKED_FILL = PatternFill("solid", fgColor="EDEDED")

BASE_HEADERS = ["Code", "Sub-component", "Indicator", "Unit", "How to report this figure", "Target"]
TAIL_HEADERS = ["Data source", "Comments"]
BASE_WIDTHS = [12, 13, 52, 10, 46, 13]
TAIL_WIDTHS = [26, 32]


def reporting_basis(indicator: Indicator, parts: list[str] | None = None) -> str:
    """Plain instruction for what figure this row wants.

    Stated per row because it differs per indicator, which is exactly what a
    single column header cannot express -- and the ambiguity the NPCU's notes
    record as "some states enter cumulative, some marginal, some snapshot".

    ``parts`` narrows a composite to the components the state actually
    implements, so Ekiti is not told its C1.0-01 must equal a sum that includes
    a sub-component 1.1 row locked out of its own template.
    """
    if indicator.unit == IndicatorUnit.BOOLEAN:
        return "Enter Yes or No."
    components = indicator.composite_of if parts is None else parts
    if components:
        return f"Cumulative total to date. Must equal {' + '.join(components)}."
    if indicator.unit == IndicatorUnit.PERCENT:
        return "Rate at the end of this period, between 0 and 100. Do not enter a count."
    if indicator.is_cumulative:
        return "Cumulative total to date, including everything reported in earlier periods."
    return "Value as at the end of this period. A snapshot, not a running total."


def _state_targets(
    db: Session, state: State, period: ReportingPeriod
) -> dict[int, float]:
    """``{indicator_id: target}`` for this state's approved targets only.

    National targets are deliberately not served here. A state shown the
    national figure reports against it, which is how a single state comes to
    claim most of a national total. A target the governance body has not yet
    cleared for the state shows as a dash instead.
    """
    rows = db.scalars(
        select(Target).where(
            Target.period_id == period.id,
            Target.state_id == state.id,
            Target.status == str(TargetStatus.APPROVED),
        )
    )
    return {
        target.indicator_id: target.target_value
        for target in rows
        if target.target_value is not None
    }


def _applicable_subcomponents(db: Session, state: State) -> set[int] | None:
    """Sub-component ids this state implements, or ``None`` when unconfigured.

    ``None`` is not the same as the empty set: it means no applicability matrix
    has been recorded for the state, in which case nothing is locked out and the
    state is served the full catalogue.
    """
    rows = list(
        db.execute(
            select(StateSubcomponent.subcomponent_id, StateSubcomponent.implements).where(
                StateSubcomponent.state_id == state.id
            )
        )
    )
    if not rows:
        return None
    return {subcomponent_id for subcomponent_id, implements in rows if implements}


def _is_blocked(indicator: Indicator, applicable: set[int] | None) -> bool:
    """True when this state does not implement the indicator's sub-component."""
    subcomponent = indicator.subcomponent
    return (
        applicable is not None
        and subcomponent is not None
        and subcomponent.id not in applicable
    )


def _prior_figures(
    db: Session, state: State, periods: list[ReportingPeriod]
) -> dict[tuple[int, int], float]:
    """``{(period_id, indicator_id): value}`` from what the state already reported."""
    figures: dict[tuple[int, int], float] = {}
    for period in periods:
        submission = db.scalar(
            select(Submission).where(
                Submission.state_id == state.id,
                Submission.period_id == period.id,
                Submission.is_current.is_(True),
            )
        )
        if submission is None:
            continue
        for value in submission.values:
            effective = value.effective_value
            if effective is not None:
                figures[(period.id, value.indicator_id)] = effective
    return figures


def build_reporting_template(
    db: Session,
    state: State,
    period: ReportingPeriod,
    *,
    prior_periods: int = 2,
) -> bytes:
    """Build the template for one state and one reporting period."""
    indicators = [
        indicator
        for indicator in reference.active_indicators(db)
        if indicator.is_reported
    ]
    applicable = _applicable_subcomponents(db, state)
    blocked_codes = {
        indicator.code
        for indicator in indicators
        if _is_blocked(indicator, applicable)
    }
    targets = _state_targets(db, state, period)

    history = reference.ordered_periods(db, period.period_type)
    earlier = [p for p in history if p.sort_key < period.sort_key][-prior_periods:]
    figures = _prior_figures(db, state, earlier)
    # A reference column the state never reported into is noise; drop it rather
    # than show a column of dashes beside the one column it has to fill.
    earlier = [p for p in earlier if any(key[0] == p.id for key in figures)]

    is_monthly = period.period_type == PeriodType.MONTHLY
    stream = "Performance tracker" if is_monthly else "Results framework"

    # Column names are load-bearing: the mapper resolves them back to canonical
    # fields on upload. The entry column has to read as a value column and the
    # reference columns beside it must not, or a state's own history would be
    # ingested as this period's figure.
    headers = (
        BASE_HEADERS
        + [f"{p.label} — as reported" for p in earlier]
        + [f"Reported value ({period.code})"]
        + TAIL_HEADERS
    )
    widths = BASE_WIDTHS + [16] * len(earlier) + [18] + TAIL_WIDTHS

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporting"

    sheet["A1"] = f"AGILE {stream} — {state.name} — {period.label}"
    sheet["A1"].font = Font(size=14, bold=True)
    sheet["A2"] = (
        f"Submission due {period.due_date.isoformat()}. Enter a figure for every row that is "
        "open. Grey rows are locked: either they do not apply to this state, or the platform "
        "supplies the value. Do not add, delete or reorder rows."
    )
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A3"] = (
        f"State: {state.name}  ·  Cohort: {state.cohort.name if state.cohort else '—'}"
        f"  ·  Period: {period.code}"
    )
    sheet["A3"].font = Font(italic=True, color="5F7080")
    sheet["A4"] = (
        "The state is confirmed when you upload, not read from this file, so a copy of "
        "another state's template cannot submit under the wrong name."
    )
    sheet["A4"].font = Font(italic=True, size=9, color="5F7080")
    for row in (2, 4):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    sheet.row_dimensions[2].height = 34

    header_row = 6
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=4)

    value_column = len(BASE_HEADERS) + len(earlier) + 1
    current_component = None
    cursor = header_row + 1
    open_rows = blocked_rows = 0

    for indicator in indicators:
        component = indicator.category.name if indicator.category else "Other"
        if component != current_component:
            current_component = component
            banner = sheet.cell(row=cursor, column=1, value=component)
            banner.font = Font(bold=True)
            banner.fill = BANNER_FILL
            sheet.merge_cells(
                start_row=cursor, start_column=1, end_row=cursor, end_column=len(headers)
            )
            cursor += 1

        subcomponent = indicator.subcomponent
        blocked = _is_blocked(indicator, applicable)

        sheet.cell(row=cursor, column=1, value=indicator.code).font = Font(bold=True)
        sheet.cell(row=cursor, column=2, value=subcomponent.code if subcomponent else "")
        name_cell = sheet.cell(row=cursor, column=3, value=indicator.name)
        name_cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=cursor, column=4, value=_unit_label(indicator))

        if blocked:
            basis = (
                f"Not applicable — {state.name} does not implement "
                f"{subcomponent.code if subcomponent else 'this sub-component'}. Leave blank."
            )
        else:
            parts = [
                code for code in (indicator.composite_of or []) if code not in blocked_codes
            ]
            basis = reporting_basis(indicator, parts)
        basis_cell = sheet.cell(row=cursor, column=5, value=basis)
        basis_cell.alignment = Alignment(wrap_text=True, vertical="top")
        basis_cell.font = Font(size=9, italic=blocked)

        target = None if blocked else targets.get(indicator.id)
        sheet.cell(row=cursor, column=6, value=target if target is not None else "—")

        for offset, prior in enumerate(earlier):
            value = None if blocked else figures.get((prior.id, indicator.id))
            sheet.cell(row=cursor, column=len(BASE_HEADERS) + 1 + offset,
                       value=value if value is not None else "—")

        entry = sheet.cell(row=cursor, column=value_column)
        for column in range(1, len(headers) + 1):
            cell = sheet.cell(row=cursor, column=column)
            if blocked:
                cell.fill = BLOCKED_FILL
                cell.protection = Protection(locked=True)
            elif column in {value_column, value_column + 1, value_column + 2}:
                cell.fill = ENTRY_FILL
                cell.protection = Protection(locked=False)
            else:
                cell.fill = LOCKED_FILL
                cell.protection = Protection(locked=True)

        if blocked:
            entry.value = "n/a"
            blocked_rows += 1
        else:
            open_rows += 1

        sheet.row_dimensions[cursor].height = 28
        cursor += 1

    # Locked cells only bite once the sheet is protected; no password, so a
    # state can unlock it if it must, and the platform validates regardless.
    sheet.protection.sheet = True
    sheet.protection.formatCells = False

    _add_guidance(workbook, state, period, stream, open_rows, blocked_rows)
    _add_reference(workbook, indicators)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _unit_label(indicator: Indicator) -> str:
    return {
        IndicatorUnit.NUMBER: "Number",
        IndicatorUnit.PERCENT: "%",
        IndicatorUnit.BOOLEAN: "Yes/No",
        IndicatorUnit.RATIO: "Ratio",
        IndicatorUnit.SCORE: "Score",
        IndicatorUnit.CURRENCY_NGN_M: "NGN m",
    }.get(IndicatorUnit(indicator.unit), indicator.unit)


def _add_guidance(
    workbook: Workbook,
    state: State,
    period: ReportingPeriod,
    stream: str,
    open_rows: int,
    blocked_rows: int,
) -> None:
    sheet = workbook.create_sheet("How to complete")
    sheet["A1"] = f"Completing the {stream.lower()} return"
    sheet["A1"].font = Font(size=13, bold=True)

    lines = [
        f"{state.name} · {period.label} · {open_rows} figures to report"
        + (f", {blocked_rows} rows locked as not applicable" if blocked_rows else ""),
        "",
        "1. Read the 'How to report this figure' column before entering anything. It differs "
        "per row: some indicators want a cumulative total to date, some a snapshot at the end "
        "of the period, some a rate, some Yes or No.",
        "2. Do not add, delete or reorder rows, and do not edit the code or indicator name. "
        "The platform matches on the code in column A.",
        "3. Grey rows are locked. Either the sub-component does not apply to this state, or "
        "the platform supplies the value. Leave them alone.",
        "4. The target column shows this state's own target, and only once the project's "
        "governance body has cleared it. A dash means no target has been approved for this "
        "state and period yet -- report what happened, not what a national figure implies.",
        "5. Earlier periods are shown for reference, taken from what this state already "
        "reported. If one of those figures is wrong, do not change it here -- it is corrected "
        "through a correction sheet, with evidence.",
        "6. Percentages are entered as a number between 0 and 100, not as a fraction and not "
        "as a count.",
        "7. Cumulative figures must never be lower than the figure reported in an earlier "
        "period. If a figure genuinely needs to come down, report it and expect a query; "
        "you will be asked for the evidence behind the change.",
        "8. Record the means of verification in 'Data source'.",
        "",
        "Leaving a figure blank is not the same as reporting zero. Blank means not reported; "
        "zero means the activity happened and the count is nil.",
    ]
    for offset, line in enumerate(lines, start=3):
        cell = sheet.cell(row=offset, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 104


def _add_reference(workbook: Workbook, indicators: list[Indicator]) -> None:
    sheet = workbook.create_sheet("Indicator reference")
    sheet.append(
        ["Code", "Previous code", "Sub-component", "Indicator", "Unit", "Basis",
         "National aggregation"]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    aggregation = {
        AggregationMethod.SUM: "Summed across states",
        AggregationMethod.AVERAGE_NONZERO: "Unweighted average of non-zero states",
        AggregationMethod.COUNT_YES: "Count of states answering Yes",
        AggregationMethod.WEIGHTED_AVERAGE: "Weighted by numerator and denominator",
        AggregationMethod.AVERAGE: "Unweighted average",
    }
    for indicator in indicators:
        sheet.append([
            indicator.code,
            indicator.legacy_code or "",
            indicator.subcomponent.code if indicator.subcomponent else "",
            indicator.name,
            _unit_label(indicator),
            reporting_basis(indicator),
            aggregation.get(
                AggregationMethod(indicator.aggregation_method), indicator.aggregation_method
            ),
        ])
    for column, width in zip("ABCDEFG", [12, 13, 14, 56, 10, 52, 38], strict=True):
        sheet.column_dimensions[column].width = width
