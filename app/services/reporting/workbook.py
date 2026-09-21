"""The analysis model as the workbook the NPCU already works in.

Every quarter the NPCU builds an analysis workbook by hand -- one sheet per
component, every state a column, the national total and the target beside it,
and the flagged cells shaded. It is the artefact the Bank sees and the one the
states are shown their own figures in. Rebuilding it by hand each cycle is
where the transcription errors come from, and it is the step this replaces.

The sheets are the ones that workbook already has, in the order it has them,
so a reader opening this one needs no explanation. The figures are the same
arrays the dashboard draws from -- not a second calculation that could drift
from the first.

**Flagged figures are shaded, never dropped.** Holding the doubtful ones out
of the totals is what made this platform's national figures disagree with the
NPCU's own published technical report, by nearly 200,000 on one indicator. A
flagged cell carries its value, its shading and its note, so a reader sees
both what it contributes and what it distorts.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.analysis_model import (
    COMPONENT_NAMES,
    COMPONENT_SHORT,
    AnalysisModel,
    IndicatorRow,
)

#: The NPCU's own palette, so the two workbooks sit side by side.
NAVY = "1F4E79"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(color=NAVY, bold=True, size=14)
LABEL_FONT = Font(bold=True, size=10)
NOTE_FONT = Font(italic=True, size=9, color="595959")

NATIONAL_FILL = PatternFill("solid", fgColor="BDD7EE")
TARGET_FILL = PatternFill("solid", fgColor="FFF2CC")

#: Flag shading, matching the dashboard band for band.
FLAG_FILL = {
    "C": PatternFill("solid", fgColor="F4B7B7"),
    "H": PatternFill("solid", fgColor="F8CBAD"),
    "M": PatternFill("solid", fgColor="FFE699"),
}
FLAG_FONT = {
    "C": Font(color="9C0006", bold=True, size=10),
    "H": Font(color="974706", bold=True, size=10),
    "M": Font(color="7F6000", bold=True, size=10),
}
FLAG_NAME = {"C": "Critical", "H": "High", "M": "Medium"}

ACHIEVEMENT_FILL = {
    "good": PatternFill("solid", fgColor="C6EFCE"),
    "warn": PatternFill("solid", fgColor="FFEB9C"),
    "bad": PatternFill("solid", fgColor="FFC7CE"),
}

THIN = Side(style="thin", color="D6DEE8")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _band(pct: float | None) -> str | None:
    if pct is None:
        return None
    return "good" if pct >= 80 else "warn" if pct >= 50 else "bad"


def _number_format(row: IndicatorRow) -> str:
    if row.is_boolean:
        return "General"
    return '#,##0.0"%"' if row.unit == "PERCENT" else "#,##0.0" if row.is_rate else "#,##0"


def _header(sheet: Worksheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = CELL_BORDER
    sheet.row_dimensions[row].height = 34
    sheet.freeze_panes = sheet.cell(row=row + 1, column=4)


def _title(sheet: Worksheet, text: str, subtitle: str | None = None) -> int:
    sheet.cell(row=1, column=1, value=text).font = TITLE_FONT
    if subtitle:
        sheet.cell(row=2, column=1, value=subtitle).font = NOTE_FONT
        return 4
    return 3


# --------------------------------------------------------------------------
# Sheets
# --------------------------------------------------------------------------
def _cover(book: Workbook, model: AnalysisModel) -> None:
    sheet = book.create_sheet("Cover")
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 86

    sheet.cell(row=1, column=1, value="AGILE Project — Quarterly Analysis Model").font = (
        TITLE_FONT
    )
    facts = [
        ("Project", "P170664 — Adolescent Girls Initiative for Learning and Empowerment"),
        ("Implementing agency", "Federal Ministry of Education, Nigeria"),
        ("Reporting period", model.period_label),
        ("States reporting", f"{model.states_reporting} of {len(model.states)}"),
        ("Indicators", f"{len(model.indicators)} reported indicators"),
        ("Open findings", f"{len(model.flags)} finding kinds across the period's returns"),
        ("Source", "KoBoToolbox state submissions, validated in the AGILE reporting platform"),
        ("Generated", datetime.now(timezone.utc).strftime("%d %B %Y, %H:%M UTC")),
    ]
    row = 3
    for label, value in facts:
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        sheet.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="How to read the flags").font = LABEL_FONT
    row += 1
    for note in (
        "Flagged figures are shaded, not removed. Every figure a state reported "
        "counts towards the national total whatever its flag.",
        "Excluding the doubtful figures made this platform's totals disagree with "
        "the NPCU's own published technical report, so a flagged cell now carries "
        "its value, its shading and its note.",
        "A blank cell is an unreported figure, which is not the same as a reported "
        "zero. Rate indicators are the unweighted average of the states that reported.",
        "A figure changes only through the query workflow, never because a rule "
        "decided it should.",
    ):
        cell = sheet.cell(row=row, column=2, value=note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 30
        row += 1

    row += 1
    for band, label in (
        ("C", "Critical — the figure cannot be true as reported"),
        ("H", "High — movement large enough to change a national reading"),
        ("M", "Medium — a coverage gap or an unexplained change"),
    ):
        marker = sheet.cell(row=row, column=1, value=FLAG_NAME[band])
        marker.fill = FLAG_FILL[band]
        marker.font = FLAG_FONT[band]
        marker.border = CELL_BORDER
        sheet.cell(row=row, column=2, value=label)
        row += 1


def _national_summary(book: Workbook, model: AnalysisModel) -> None:
    sheet = book.create_sheet("National Summary")
    start = _title(
        sheet,
        f"National summary — {model.period_label}",
        "Every reported indicator, its national position and whether it carries a flag.",
    )
    _header(
        sheet,
        start,
        [
            "Code", "Indicator", "Component", "Type", "National", "Target",
            "% of target", "States reporting", "Flag", "Note",
        ],
    )

    row = start + 1
    for indicator in model.indicators:
        sheet.cell(row=row, column=1, value=indicator.code).font = LABEL_FONT
        sheet.cell(row=row, column=2, value=indicator.name).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=3, value=indicator.component)
        sheet.cell(row=row, column=4, value=indicator.type_label)

        national = sheet.cell(row=row, column=5, value=indicator.achieved)
        national.number_format = _number_format(indicator)
        national.fill = (
            FLAG_FILL[indicator.national_flag] if indicator.national_flag else NATIONAL_FILL
        )

        target = sheet.cell(row=row, column=6, value=indicator.target)
        target.number_format = _number_format(indicator)
        target.fill = TARGET_FILL

        pct = sheet.cell(row=row, column=7, value=indicator.achievement_pct)
        pct.number_format = '0.0"%"'
        band = _band(indicator.achievement_pct)
        if band:
            pct.fill = ACHIEVEMENT_FILL[band]

        sheet.cell(
            row=row, column=8, value=f"{indicator.reporting} of {indicator.expected}"
        ).alignment = Alignment(horizontal="center")

        flagged = indicator.national_flag or (
            max(indicator.flag_severity.values(), key=lambda s: -"CHM".index(s))
            if indicator.flag_severity
            else None
        )
        flag_cell = sheet.cell(
            row=row, column=9, value=FLAG_NAME[flagged] if flagged else ""
        )
        if flagged:
            flag_cell.fill = FLAG_FILL[flagged]
            flag_cell.font = FLAG_FONT[flagged]

        sheet.cell(row=row, column=10, value=indicator.flag_note).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        row += 1

    widths = {1: 12, 2: 56, 3: 12, 4: 16, 5: 15, 6: 14, 7: 12, 8: 16, 9: 11, 10: 70}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def _component_sheet(book: Workbook, model: AnalysisModel, component: str) -> None:
    rows = model.by_component(component)
    if not rows:
        return
    # Excel caps a sheet name at 31 characters, which cuts the full component
    # title mid-word ("Component 1 — Creating Safe and"). The tab carries the
    # short name the NPCU's own workbook uses; the full title is the heading.
    sheet = book.create_sheet(COMPONENT_SHORT.get(component, component)[:31])
    start = _title(
        sheet,
        COMPONENT_NAMES.get(component, component),
        f"{model.period_label} — every state's figure, with the national total and target.",
    )
    _header(
        sheet,
        start,
        ["Code", "Indicator", "Type"] + list(model.states)
        + ["National Total", "Target", "% of Target"],
    )

    state_start = 4
    row = start + 1
    for indicator in rows:
        sheet.cell(row=row, column=1, value=indicator.code).font = LABEL_FONT
        name = sheet.cell(row=row, column=2, value=indicator.name)
        name.alignment = Alignment(wrap_text=True, vertical="top")
        if indicator.flag_note:
            name.comment = _note(indicator.flag_note)
        sheet.cell(row=row, column=3, value=indicator.type_label)

        for offset, state in enumerate(model.states):
            value = indicator.states[offset]
            cell = sheet.cell(row=row, column=state_start + offset)
            # A boolean reads Yes or No; a blank stays blank, because an
            # unreported figure is not a reported zero.
            if value is None:
                cell.value = None
            elif indicator.is_boolean:
                cell.value = "Yes" if value >= 1 else "No"
            else:
                cell.value = value
                cell.number_format = _number_format(indicator)
            severity = indicator.flag_severity.get(state)
            if severity:
                cell.fill = FLAG_FILL[severity]
                cell.font = FLAG_FONT[severity]
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="right")

        national_col = state_start + len(model.states)
        national = sheet.cell(row=row, column=national_col, value=indicator.achieved)
        national.number_format = _number_format(indicator)
        national.font = LABEL_FONT
        national.fill = (
            FLAG_FILL[indicator.national_flag] if indicator.national_flag else NATIONAL_FILL
        )

        target = sheet.cell(row=row, column=national_col + 1, value=indicator.target)
        target.number_format = _number_format(indicator)
        target.fill = TARGET_FILL

        pct = sheet.cell(
            row=row, column=national_col + 2, value=indicator.achievement_pct
        )
        pct.number_format = '0.0"%"'
        band = _band(indicator.achievement_pct)
        if band:
            pct.fill = ACHIEVEMENT_FILL[band]
        row += 1

    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 52
    sheet.column_dimensions["C"].width = 16
    for offset in range(len(model.states)):
        sheet.column_dimensions[get_column_letter(state_start + offset)].width = 13
    for offset in range(3):
        sheet.column_dimensions[
            get_column_letter(state_start + len(model.states) + offset)
        ].width = 15


def _note(text: str):
    from openpyxl.comments import Comment

    comment = Comment(text, "AGILE reporting platform")
    comment.width = 420
    comment.height = 130
    return comment


def _flags_sheet(book: Workbook, model: AnalysisModel) -> None:
    sheet = book.create_sheet("Data Quality Flags")
    start = _title(
        sheet,
        f"Data quality flags — {model.period_label}",
        "Every open finding against this period's returns. Each one is also a query "
        "raised with the state that reported it.",
    )
    _header(sheet, start, ["Severity", "Code", "Indicator", "States", "Finding"])

    row = start + 1
    for flag in model.flags:
        band = "C" if flag.severity == "CRITICAL" else "H"
        severity = sheet.cell(row=row, column=1, value=flag.severity)
        severity.fill = FLAG_FILL[band]
        severity.font = FLAG_FONT[band]
        severity.border = CELL_BORDER
        sheet.cell(row=row, column=2, value=flag.code).font = LABEL_FONT
        sheet.cell(row=row, column=3, value=flag.indicator).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=4, value=flag.states_label).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        sheet.cell(row=row, column=5, value=flag.issue).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
        row += 1

        # Each state's own wording, where they differ. One state's figures
        # printed above a list of four states' names is a register nobody can
        # act on.
        for state, message in flag.details:
            sheet.cell(row=row, column=4, value=state).font = NOTE_FONT
            detail = sheet.cell(row=row, column=5, value=message)
            detail.font = NOTE_FONT
            detail.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            row += 1

    for column, width in {1: 12, 2: 13, 3: 46, 4: 30, 5: 96}.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = sheet.cell(row=start + 1, column=1)


def _full_table(book: Workbook, model: AnalysisModel) -> None:
    """Every indicator and every state on one sheet, for filtering and pivots."""
    sheet = book.create_sheet("Full Data Table")
    _header(
        sheet,
        1,
        ["Component", "Code", "Indicator", "Type", "State", "Value", "Flag", "Finding"],
    )
    row = 2
    for indicator in model.indicators:
        for offset, state in enumerate(model.states):
            value = indicator.states[offset]
            if value is None:
                continue  # an unreported figure is a gap, not a row of zero
            severity = indicator.flag_severity.get(state)
            sheet.cell(row=row, column=1, value=indicator.component)
            sheet.cell(row=row, column=2, value=indicator.code)
            sheet.cell(row=row, column=3, value=indicator.name)
            sheet.cell(row=row, column=4, value=indicator.type_label)
            sheet.cell(row=row, column=5, value=state)
            cell = sheet.cell(row=row, column=6, value=value)
            cell.number_format = _number_format(indicator)
            sheet.cell(row=row, column=7, value=FLAG_NAME[severity] if severity else "")
            sheet.cell(row=row, column=8, value=indicator.flag_note if severity else "")
            row += 1
    sheet.auto_filter.ref = f"A1:H{max(row - 1, 2)}"
    for column, width in {1: 12, 2: 12, 3: 52, 4: 16, 5: 15, 6: 15, 7: 11, 8: 70}.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------
def build(model: AnalysisModel) -> Workbook:
    """Assemble the analysis workbook for one reporting period."""
    book = Workbook()
    book.remove(book.active)

    _cover(book, model)
    _national_summary(book, model)
    for component in model.components:
        _component_sheet(book, model, component)
    _flags_sheet(book, model)
    _full_table(book, model)
    return book


def write(model: AnalysisModel, directory: Path) -> Path:
    """Write the workbook and return where it landed."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"AGILE_{model.period_code}_Analysis_Model.xlsx"
    build(model).save(path)
    return path
